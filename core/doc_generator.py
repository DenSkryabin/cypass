import io
import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    """Установка фонового цвета ячейки таблицы."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def generate_crmd_statements_docx(applicant_name, arc_nr, mp_nr, submission_date, trips_list):
    """
    Генерация официального комплекта Statement No. 1 и Statement No. 2
    по форме CRMD (Миграционный департамент Кипра).
    """
    doc = Document()

    # Поля страницы 2 см
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # -------------------------------------------------------------
    # STATEMENT NO. 1
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    run_t1 = title_p.add_run("STATEMENT NO. 1\n")
    run_t1.bold = True
    run_t1.font.size = Pt(12)
    
    info_p = doc.add_paragraph()
    info_p.paragraph_format.space_after = Pt(8)
    info_p.add_run(f"NAME OF APPLICANT: {applicant_name.upper()}\n").bold = True
    info_p.add_run(f"Date of application M127: {submission_date.strftime('%d/%m/%Y')} / MP: {mp_nr} / ARC nr.: {arc_nr}\n")
    info_p.add_run("ΚΑΤΑΣΤΑΣΗ ΠΑΡΑΜΟΝΗΣ ΣΤΗΝ ΚΥΠΡΟ ΣΥΜΦΩΝΑ ΜΕ ΤΟ ΔΙΑΒΑΤΗΡΙΟ\n").bold = True
    info_p.add_run("STATEMENT OF STAY IN CYPRUS ACCORDING TO PASSPORT").bold = True

    table1 = doc.add_table(rows=2, cols=5)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_row0 = table1.rows[0].cells
    hdr_row0[0].merge(hdr_row0[1])
    hdr_row0[0].text = "ΑΦΙΞΕΙΣ στην Κύπρο\nARRIVALS to Cyprus"
    hdr_row0[2].merge(hdr_row0[3])
    hdr_row0[2].text = "ΑΝΑΧΩΡΗΣΕΙΣ από την Κύπρο\nDEPARTURES from Cyprus"
    hdr_row0[4].text = "ΔΙΑΡΚΕΙΑ ΠΑΡΑΜΟΝΗΣ\nDURATION OF STAY"

    sub_headers = ["DATE", "PAGE", "DATE", "PAGE", "DAYS"]
    hdr_row1 = table1.rows[1].cells
    for i, h in enumerate(sub_headers):
        hdr_row1[i].text = h

    for row in table1.rows[:2]:
        for cell in row.cells:
            set_cell_background(cell, "F2F2F2")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(9)

    total_days_stay = 0
    clean_trips = [t for t in trips_list if t.get("Arrival") and t.get("Departure")]
    clean_trips.sort(key=lambda x: x["Arrival"])

    for idx, tr in enumerate(clean_trips):
        row_cells = table1.add_row().cells
        arr_d = tr["Arrival"]
        dep_d = tr["Departure"]
        
        # Последняя поездка длится до дня подачи
        if idx == len(clean_trips) - 1:
            dep_str = f"Till today\n{submission_date.strftime('%d/%m/%Y')}"
            dep_page = "Not applicable"
            days_count = (submission_date - arr_d).days
        else:
            dep_str = dep_d.strftime("%d/%m/%Y")
            dep_page = tr.get("Dep_Page", "Pass No. P XX")
            # Правило Кипра: день въезда считается, день выезда — нет
            days_count = (dep_d - arr_d).days

        total_days_stay += days_count

        row_cells[0].text = arr_d.strftime("%d/%m/%Y")
        row_cells[1].text = tr.get("Arr_Page", "Pass No. P XX")
        row_cells[2].text = dep_str
        row_cells[3].text = dep_page
        row_cells[4].text = str(days_count)

        for i in range(5):
            row_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in row_cells[i].paragraphs[0].runs:
                r.font.size = Pt(9)

    tot_row = table1.add_row().cells
    tot_row[0].merge(tot_row[3])
    tot_row[0].text = "ΣΥΝΟΛΟ ΠΑΡΑΜΟΝΗΣ / TOTAL STAY"
    tot_row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tot_row[4].text = str(total_days_stay)
    tot_row[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in tot_row[0].paragraphs[0].runs + tot_row[4].paragraphs[0].runs:
        r.bold = True
        r.font.size = Pt(10)

    decl_p = doc.add_paragraph()
    decl_p.paragraph_format.space_before = Pt(12)
    decl_p.add_run(f"DECLARATION: I the undersign {applicant_name.upper()}\n").bold = True
    decl_p.add_run("Do solemnly declare that the foregoing particulars stated above are true and I make this solemn declaration knowing the provisions of the Law for untrue declaration.\n")
    decl_p.add_run(f"Date: {submission_date.strftime('%d/%m/%Y')}        Signature: .....................................................")

    doc.add_page_break()

    # -------------------------------------------------------------
    # STATEMENT NO. 2 (Финальные 12 месяцев)
    # -------------------------------------------------------------
    p2 = doc.add_paragraph()
    r2 = p2.add_run("STATEMENT NO. 2\n")
    r2.bold = True
    r2.font.size = Pt(12)

    p2_info = doc.add_paragraph()
    p2_info.add_run(f"NAME OF APPLICANT: {applicant_name.upper()}\n").bold = True
    p2_info.add_run(f"Date of application M127: {submission_date.strftime('%d/%m/%Y')} / MP: {mp_nr} / ARC nr.: {arc_nr}\n")
    p2_info.add_run("ΚΑΤΑΣΤΑΣΗ ΠΑΡΑΜΟΝΗΣ ΑΛΛΟΔΑΠΟΥ ΤΟΥΣ ΤΕΛΕΥΤΑΙΟΥΣ 12 ΜΗΝΕΣ ΠΡΙΝ ΑΠΟ ΤΗΝ ΥΠΟΒΟΛΗ ΤΗΣ ΑΙΤΗΣΗΣ Μ127\n").bold = True
    p2_info.add_run("STATEMENT OF STAY IN CYPRUS DURING THE LAST 12 MONTHS BEFORE THE DATE OF APPLICATION M127").bold = True

    table2 = doc.add_table(rows=2, cols=5)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    h2_row0 = table2.rows[0].cells
    h2_row0[0].merge(h2_row0[1])
    h2_row0[0].text = "ΑΦΙΞΕΙΣ στην Κύπρο\nARRIVALS to Cyprus"
    h2_row0[2].merge(h2_row0[3])
    h2_row0[2].text = "ΑΝΑΧΩΡΗΣΕΙΣ από την Κύπρο\nDEPARTURES from Cyprus"
    h2_row0[4].text = "ΔΙΑΡΚΕΙΑ ΠΑΡΑΜΟΝΗΣ\nDURATION OF STAY"

    h2_row1 = table2.rows[1].cells
    for i, h in enumerate(sub_headers):
        h2_row1[i].text = h

    for row in table2.rows[:2]:
        for cell in row.cells:
            set_cell_background(cell, "F2F2F2")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(9)

    cutoff_12m = submission_date - datetime.timedelta(days=365)
    last_year_trips = [t for t in clean_trips if t["Departure"] > cutoff_12m]

    tot_12m = 0
    for idx, tr in enumerate(last_year_trips):
        row_c = table2.add_row().cells
        arr_d = max(tr["Arrival"], cutoff_12m)
        dep_d = tr["Departure"]
        
        is_last = (idx == len(last_year_trips) - 1)
        if is_last:
            dep_str = f"Till today\n{submission_date.strftime('%d/%m/%Y')}"
            dep_page = "Not applicable"
            days = (submission_date - arr_d).days
        else:
            dep_str = dep_d.strftime("%d/%m/%Y")
            dep_page = tr.get("Dep_Page", "Pass No. P XX")
            days = (dep_d - arr_d).days

        tot_12m += days
        row_c[0].text = arr_d.strftime("%d/%m/%Y") + (" (αναλογία)" if arr_d == cutoff_12m else "")
        row_c[1].text = tr.get("Arr_Page", "Pass No. P XX")
        row_c[2].text = dep_str
        row_c[3].text = dep_page
        row_c[4].text = str(days)

        for i in range(5):
            row_c[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in row_c[i].paragraphs[0].runs:
                r.font.size = Pt(9)

    tot2_row = table2.add_row().cells
    tot2_row[0].merge(tot2_row[3])
    tot2_row[0].text = "ΣΥΝΟΛΟ ΠΑΡΑΜΟΝΗΣ / TOTAL STAY"
    tot2_row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tot2_row[4].text = str(tot_12m)
    tot2_row[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in tot2_row[0].paragraphs[0].runs + tot2_row[4].paragraphs[0].runs:
        r.bold = True
        r.font.size = Pt(10)

    decl2_p = doc.add_paragraph()
    decl2_p.paragraph_format.space_before = Pt(12)
    decl2_p.add_run(f"DECLARATION: I the undersign {applicant_name.upper()}\n").bold = True
    decl2_p.add_run("Do solemnly declare that the foregoing particulars stated above are true and I make this solemn declaration knowing the provisions of the Law for untrue declaration.\n")
    decl2_p.add_run(f"Date: {submission_date.strftime('%d/%m/%Y')}        Signature: .....................................................")

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output